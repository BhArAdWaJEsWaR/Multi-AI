
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9.5)

# Header
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A BHARADWAJ ESWAR")
r.bold = True
r.font.size = Pt(17)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("bharadwajavvaru5@gmail.com  |  +91-9441585019  |  LinkedIn  |  GitHub")
r.font.size = Pt(9)

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11.5)
    return p

def add_entry(title, details, bold_details=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(title)
    r.bold = True
    if details:
        r2 = p.add_run("  " + details)
        if bold_details:
            r2.bold = True
    return p

heading("EDUCATION")

add_entry("Amrita Vishwa Vidyapeetham, Coimbatore", "")
p = doc.add_paragraph("B.Tech in Computer Science and Engineering (AI)  |  2023 – 2027")
p.paragraph_format.space_after = Pt(1)
p.add_run("CGPA: 7.3/10").bold = True

add_entry("Sasi Educational Institutes, Velivennu", "")
p = doc.add_paragraph("Higher Secondary  |  2021 – 2023")
p.paragraph_format.space_after = Pt(1)
p.add_run("Percentage: 94.5%").bold = True

add_entry("Narayana High School", "")
p = doc.add_paragraph("Matriculation  |  2021")
p.paragraph_format.space_after = Pt(1)
p.add_run("Grade: 596/600").bold = True

heading("SKILLS")

skills = [
    ("Programming", "Python, C++, C, SQL"),
    ("AI/ML", "LLMs, Deep Learning, CNN, LSTM, Computer Vision, Embeddings, Semantic Classification"),
    ("Frameworks & Libraries", "TensorFlow/Keras, OpenCV, NumPy, SymPy, MediaPipe, YOLO"),
    ("AI Systems", "Ollama, Local LLM Inference, Multi-Agent Systems, Model Routing, REST APIs"),
    ("Robotics & Tools", "Embedded Systems, Motor Control, Autonomous Navigation, Git, GitHub, VS Code, MATLAB, Render"),
]
for label, value in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.add_run(label + ": ").bold = True
    p.add_run(value)

heading("CERTIFICATIONS")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
p.add_run("The Complete Python Bootcamp From Zero to Hero in Python").bold = True
p = doc.add_paragraph("Udemy  |  Issued: September 2026")
p.paragraph_format.space_after = Pt(1)
p = doc.add_paragraph("Credential ID: UC-31ddb6e2-df09-48e5-98a0-c73267ad1708")
p.paragraph_format.space_after = Pt(1)
p = doc.add_paragraph("Skills: Python Programming")
p.paragraph_format.space_after = Pt(1)

heading("ACADEMIC PROJECTS")

projects = [
    ("1. Multi-Agent AI Orchestration System", [
        "Designed and developed a resource-efficient multi-agent AI orchestration system using locally hosted open-source LLMs through Ollama.",
        "Implemented a hierarchical model cascade that routes tasks from lightweight models to specialized models and escalates complex or uncertain queries to larger models.",
        "Built embedding-based semantic classification for specialist selection, along with validation and escalation pipelines for response evaluation.",
        "Implemented dynamic model loading and unloading to reduce RAM consumption while maintaining conversation context for efficient local inference on memory-constrained hardware.",
        "Technologies: Python, Ollama, LLMs, Embeddings, Semantic Classification, Multi-Agent Systems, Model Routing, Local Inference"
    ]),
    ("2. Human-Following Robot", [
        "Designed a human-tracking mobile robot using two vision approaches, YOLO-based person detection and a lightweight MediaPipe pipeline, to overcome hardware constraints.",
        "Implemented distance-aware motion control enabling the robot to maintain an optimal following distance, stop when the subject is too close, and reverse when necessary.",
        "Integrated real-time inference, motor control logic, and safety checks to enable smooth autonomous navigation.",
        "Technologies: Python, YOLO, MediaPipe, OpenCV, Embedded Systems, Motor Control, Robotics"
    ]),
    ("3. Handwritten Equation Solver", [
        "Built a deep learning-based system to recognize and solve handwritten mathematical equations.",
        "Combined CNN and LSTM architectures to convert handwritten expressions into LaTeX representations and used SymPy for symbolic computation to generate solutions.",
        "Integrated OpenCV preprocessing to enhance handwriting clarity and improve model performance.",
        "Technologies: Python, TensorFlow/Keras, CNN + LSTM, SymPy, Deep Learning"
    ])
]

for title, bullets in projects:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.add_run(title).bold = True
    for bullet in bullets:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1)
        p.add_run("• ")
        p.add_run(bullet)

# Save
path = "A_Bharadwaj_Eswar_Resume.docx"
doc.save(path)
print(path)
